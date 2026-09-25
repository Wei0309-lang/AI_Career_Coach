# 履歷：讀取 / 儲存 + AI 健檢 / PDF、Word 上傳解析
import io

import pytest
from docx import Document
from docx.oxml import parse_xml

import Main
import Model
from resume_utils import flatten_resume_value

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
RESUME = {"fullName": "王小明", "summary": "資工系應屆畢業生", "skills": "Python、React", "experience": "AI 面試平台專題"}


def docx_bytes(build):
    document = Document()
    build(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def table_resume(document):
    """常見履歷範本：姓名電話在頁首、主要內容在表格、自傳在文字方塊。"""
    document.sections[0].header.paragraphs[0].text = "王小明｜0912-345-678"
    document.add_paragraph("個人履歷")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "技能", "Python、FastAPI、React、PostgreSQL"
    table.cell(1, 0).text, table.cell(1, 1).text = "經歷", "AI 模擬面試平台專題"
    table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "巢狀表格：Supabase 登入整合"
    textbox = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
               'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
               '<w:p><w:r><w:t>自傳：熱愛後端與系統設計</w:t></w:r></w:p>'
               '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
    # Word 常把文字方塊存兩份(新格式 + 相容舊版的備援)
    document.element.body.append(parse_xml(textbox))
    document.element.body.append(parse_xml(textbox))


def upload(client, content, name="cv.docx", mime=DOCX_MIME):
    return client.post("/api/resume/parse", files={"file": (name, content, mime)})


# ---------- Word 抽字 ----------
def test_docx_extraction_reads_tables_textboxes_and_headers():
    """修正前只讀內文最外層段落，表格式履歷只抽得到「個人履歷」4 個字。"""
    text = Main._extract_text_from_docx(docx_bytes(table_resume))
    for piece in ["個人履歷", "Python、FastAPI、React、PostgreSQL", "巢狀表格：Supabase 登入整合",
                  "0912-345-678", "自傳：熱愛後端與系統設計"]:
        assert piece in text
    assert text.count("自傳：熱愛後端與系統設計") == 1


def test_docx_extraction_of_plain_document_is_unchanged():
    text = Main._extract_text_from_docx(docx_bytes(lambda d: d.add_paragraph("一般段落的履歷內容，應該照常抽出。")))
    assert text == "一般段落的履歷內容，應該照常抽出。"


# ---------- 上傳解析 API ----------
def test_parse_word_resume_sends_table_text_and_requests_json(client, login, fake_gemini):
    login("u1")
    fake_gemini.reply = '{"fullName": "王小明", "summary": "應屆畢業生", "skills": ["Python", "React"], "experience": "專題"}'
    r = upload(client, docx_bytes(table_resume))
    assert r.status_code == 200
    assert r.json()["skills"] == "Python、React"  # 陣列攤平成文字(422 修正)
    call = fake_gemini.calls[-1]
    assert "Python、FastAPI、React、PostgreSQL" in call["contents"]
    assert call["config"].response_mime_type == "application/json"


@pytest.mark.parametrize("reply", [
    '```json\n{"fullName": "王小明", "summary": "", "skills": "", "experience": ""}\n```',
    '以下是整理結果：{"fullName": "王小明", "summary": "", "skills": "", "experience": ""} 希望有幫助',
])
def test_parse_tolerates_fences_and_surrounding_text(client, login, fake_gemini, reply):
    login("u1")
    fake_gemini.reply = reply
    r = upload(client, docx_bytes(table_resume))
    assert r.status_code == 200 and r.json()["fullName"] == "王小明"


@pytest.mark.parametrize("reply", ["這份文件內容不完整，無法整理。", '["不是物件"]', RuntimeError("429 RESOURCE_EXHAUSTED")])
def test_parse_returns_generic_502_when_ai_fails(client, login, fake_gemini, reply):
    login("u1")
    fake_gemini.reply = reply
    r = upload(client, docx_bytes(table_resume))
    assert r.status_code == 502
    assert r.json()["detail"] == "AI 解析履歷內容失敗，請稍後再試"


def test_parse_rejects_unsupported_extension(client, login, fake_gemini):
    login("u1")
    assert upload(client, b"hello", name="cv.txt", mime="text/plain").status_code == 400


def test_parse_rejects_oversized_file_without_calling_ai(client, login, fake_gemini):
    login("u1")
    big = b"%PDF-1.4\n" + b"0" * (Main.MAX_RESUME_UPLOAD_MB * 1024 * 1024 + 10)
    r = upload(client, big, name="big.pdf", mime="application/pdf")
    assert r.status_code == 400 and "上限" in r.json()["detail"]
    assert fake_gemini.calls == []


def test_parse_rejects_document_without_selectable_text(client, login, fake_gemini):
    login("u1")
    r = upload(client, docx_bytes(lambda d: d.add_paragraph("短")))
    assert r.status_code == 400 and "偵測不到" in r.json()["detail"]


# ---------- 讀取 / 儲存 ----------
def test_get_resume_returns_empty_strings_for_new_user(client, login):
    login("u1")
    assert client.get("/api/resume").json() == {"fullName": "", "summary": "", "skills": "", "experience": ""}


def test_submit_resume_saves_and_returns_suggestion(client, login, fake_gemini, db):
    login("u1")
    fake_gemini.reply = "1. 多寫量化成果"
    r = client.post("/api/resume", json=RESUME)
    assert r.status_code == 200 and r.json()["suggestion"] == "1. 多寫量化成果"
    assert db.get(Model.User, "u1").skills == "Python、React"
    assert client.get("/api/resume").json() == RESUME


def test_submit_resume_coerces_non_string_fields_instead_of_422(client, login, fake_gemini, db):
    """AI 解析結果偶爾是陣列/物件/null，儲存端點本身要能攤平，不能直接 422。"""
    login("u1")
    r = client.post("/api/resume", json={"fullName": None, "summary": "s", "skills": ["Python", "React"],
                                         "experience": [{"公司": "A", "職稱": "工程師"}, {"公司": "B"}]})
    assert r.status_code == 200
    user = db.get(Model.User, "u1")
    assert (user.full_name, user.skills, user.experience) == ("", "Python、React", "A、工程師\nB")


def test_submit_resume_still_saves_when_ai_fails_and_hides_error(client, login, fake_gemini, db):
    login("u1")
    fake_gemini.reply = RuntimeError("SECRET-INTERNAL-XYZ")
    r = client.post("/api/resume", json=RESUME)
    assert r.status_code == 200
    assert "已儲存" in r.json()["suggestion"] and "SECRET-INTERNAL-XYZ" not in r.text
    assert db.get(Model.User, "u1").full_name == "王小明"


@pytest.mark.parametrize("value, expected", [
    (None, ""), ("  文字  ", "文字"), (3, "3"), (["a", "", "b"], "a、b"),
    ({"x": "a", "y": "b"}, "a、b"), ([{"x": "a"}, {"y": "b"}], "a\nb"),
])
def test_flatten_resume_value(value, expected):
    assert flatten_resume_value(value) == expected
